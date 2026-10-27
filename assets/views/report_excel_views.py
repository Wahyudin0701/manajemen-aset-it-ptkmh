import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.db.models import Q

from assets.models import Asset, AssetUnit, Category, Employee, AssetAssignment, UnitLog

@login_required
def asset_export_excel(request):
    """Generate Excel report for assets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    from django.db.models import Count, Q
    import datetime
    
    # Get the tab filter
    tab = request.GET.get('tab', 'all')
    search_query = request.GET.get('q', '').strip()
    
    assets = Asset.objects.annotate(
        tersedia_count=Count('units', filter=Q(units__status='tersedia')),
        digunakan_count=Count('units', filter=Q(units__status='digunakan')),
        rusak_count=Count('units', filter=Q(units__status='rusak')),
        total_unit_count=Count('units')
    ).order_by('name')
    
    category_label = 'Semua Kategori'
    
    if tab == 'karyawan':
        assets = assets.exclude(category__name__icontains='Network')
        category_label = 'Aset Karyawan'
    elif tab == 'networking':
        assets = assets.filter(category__name__icontains='Network')
        category_label = 'Aset Network & Infrastruktur'
        
    if search_query:
        assets = assets.filter(
            Q(name__icontains=search_query) | 
            Q(asset_code__icontains=search_query) |
            Q(brand__icontains=search_query)
        )
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Daftar Aset"
    
    # Header styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    # Write metadata
    ws.merge_cells('A1:I1')
    ws['A1'] = "LAPORAN DAFTAR ASET IT"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    ws.merge_cells('A2:I2')
    ws['A2'] = f"Kategori: {category_label}"
    ws['A2'].alignment = center_align
    
    ws.merge_cells('A3:I3')
    ws['A3'] = f"Tanggal Dicetak: {datetime.datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws['A3'].alignment = center_align
    
    # Headers
    headers = ['No', 'Kode Aset', 'Kategori', 'Merek', 'Nama Aset', 'Total (Qty)', 'Tersedia', 'Digunakan', 'Maintenance']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    # Write data
    for idx, asset in enumerate(assets, 1):
        row_num = idx + 5
        data = [
            idx,
            asset.asset_code,
            asset.category.name if asset.category else '-',
            asset.brand,
            asset.name,
            asset.qty,
            asset.tersedia_count,
            asset.digunakan_count,
            asset.rusak_count
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 6, 7, 8, 9]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    # Adjust column widths
    column_widths = {'A': 5, 'B': 15, 'C': 20, 'D': 15, 'E': 40, 'F': 12, 'G': 12, 'H': 12, 'I': 15}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
        
    # Prepare response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Laporan_Aset_{now_str}.xlsx"'
    
    wb.save(response)
    return response



@login_required
def distribusi_export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    import datetime
    
    tab = request.GET.get('tab', 'semua')
    assignments = AssetAssignment.objects.select_related('unit__asset', 'employee', 'from_employee').order_by('-assigned_date')
    
    if tab == 'aktif':
        assignments = assignments.filter(status='aktif')
        status_label = 'Aktif'
    elif tab == 'kembali':
        assignments = assignments.filter(status='dikembalikan')
        status_label = 'Dikembalikan'
    else:
        status_label = 'Semua'
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Distribusi"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:G1')
    ws['A1'] = "LAPORAN DISTRIBUSI ASET"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    ws.merge_cells('A2:G2')
    ws['A2'] = f"Status: {status_label}"
    ws['A2'].alignment = center_align
    
    headers = ['No', 'Tanggal', 'Tipe', 'Kode Unit', 'Aset', 'Pihak Terkait', 'Status']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    for idx, a in enumerate(assignments, 1):
        row_num = idx + 4
        
        pihak = ""
        if a.type == 'handover':
            pihak = a.employee.name if a.employee else "-"
        elif a.type == 'return':
            pihak = f"Dari: {a.employee.name if a.employee else '-'}"
        elif a.type == 'mutation':
            from_name = a.from_employee.name if a.from_employee else "-"
            to_name = a.employee.name if a.employee else "-"
            pihak = f"{from_name} -> {to_name}"
            
        tipe = dict(AssetAssignment.TYPE_CHOICES).get(a.type, a.type)
        status_str = "Aktif" if a.status == 'aktif' else "Dikembalikan"
        
        data = [
            idx,
            a.assigned_date.strftime('%d %b %Y') if a.assigned_date else '-',
            tipe,
            a.unit.unit_code if a.unit else '-',
            a.unit.asset.name if a.unit and a.unit.asset else '-',
            pihak,
            status_str
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 2, 3, 4, 7]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    for col, width in zip(['A','B','C','D','E','F','G'], [5, 15, 15, 15, 30, 30, 15]):
        ws.column_dimensions[col].width = width
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Distribusi_{now_str}.xlsx"'
    wb.save(response)
    return response

@login_required
def status_aset_export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    from django.db.models import Count, Q
    import datetime
    
    tab = request.GET.get('tab', 'semua')
    
    units = AssetUnit.objects.select_related('asset', 'asset__category').all()
    if tab == 'tersedia':
        units = units.filter(status='tersedia')
        title_label = "Aset Tersedia"
    elif tab == 'digunakan':
        units = units.filter(status='digunakan')
        title_label = "Aset Digunakan"
    elif tab == 'rusak':
        units = units.filter(status='rusak')
        title_label = "Aset Maintenance / Rusak"
    else:
        title_label = "Semua Status Aset"
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Status Aset"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:F1')
    ws['A1'] = "LAPORAN STATUS UNIT ASET"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    ws.merge_cells('A2:F2')
    ws['A2'] = f"Kategori: {title_label}"
    ws['A2'].alignment = center_align
    
    headers = ['No', 'Kode Unit', 'Nama Aset', 'Area', 'Status', 'Lokasi']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    for idx, u in enumerate(units, 1):
        row_num = idx + 4
        
        status_map = {'tersedia': 'Tersedia', 'digunakan': 'Digunakan', 'rusak': 'Maintenance'}
        
        data = [
            idx,
            u.unit_code,
            u.asset.name if u.asset else '-',
            u.area or '-',
            status_map.get(u.status, u.status),
            u.location or '-'
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 2, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    for col, width in zip(['A','B','C','D','E','F'], [5, 15, 30, 20, 15, 30]):
        ws.column_dimensions[col].width = width
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Status_Aset_{now_str}.xlsx"'
    wb.save(response)
    return response

@login_required
def employee_export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    import datetime
    
    employees = Employee.objects.all().order_by('name')
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Karyawan"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:D1')
    ws['A1'] = "LAPORAN DATA KARYAWAN"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    headers = ['No', 'NIK', 'Nama Karyawan', 'Jabatan']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    for idx, e in enumerate(employees, 1):
        row_num = idx + 4
        
        data = [
            idx,
            e.nik,
            e.name,
            e.position or '-'
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 2]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    for col, width in zip(['A','B','C','D'], [5, 15, 35, 35]):
        ws.column_dimensions[col].width = width
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Karyawan_{now_str}.xlsx"'
    wb.save(response)
    return response

@login_required
def employee_detail_export_excel(request, emp_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    from django.shortcuts import get_object_or_404
    import datetime
    
    employee = get_object_or_404(Employee, id=emp_id)
    assignments = AssetAssignment.objects.filter(employee=employee, status='aktif').select_related('unit__asset')
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Aset Karyawan"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:E1')
    ws['A1'] = "LAPORAN ASET KARYAWAN"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    ws.merge_cells('A2:E2')
    ws['A2'] = f"{employee.name} ({employee.nik})"
    ws['A2'].alignment = center_align
    
    headers = ['No', 'Kode Unit', 'Nama Aset', 'Area', 'Tanggal Diberikan']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    for idx, a in enumerate(assignments, 1):
        row_num = idx + 4
        
        data = [
            idx,
            a.unit.unit_code if a.unit else '-',
            a.unit.asset.name if a.unit and a.unit.asset else '-',
            a.unit.area if a.unit else '-',
            a.assigned_date.strftime('%d %b %Y') if a.assigned_date else '-'
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 2, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    for col, width in zip(['A','B','C','D','E'], [5, 15, 30, 20, 15]):
        ws.column_dimensions[col].width = width
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Aset_{employee.name}_{now_str}.xlsx"'
    wb.save(response)
    return response

@login_required
def unit_history_export_excel(request, unit_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    from django.shortcuts import get_object_or_404
    import datetime
    
    unit = get_object_or_404(AssetUnit, id=unit_id)
    logs_qs = unit.logs.all().select_related('employee')
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Riwayat Unit"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:E1')
    ws['A1'] = "LAPORAN RIWAYAT UNIT ASET"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    ws.merge_cells('A2:E2')
    ws['A2'] = f"{unit.unit_code} - {unit.asset.name}"
    ws['A2'].alignment = center_align
    
    headers = ['No', 'Tanggal', 'Aksi', 'Keterangan', 'Pihak Terkait']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    import re
    pattern = re.compile(r"^(.+?) \((.+?)\) - (.*)$")
    
    for idx, log in enumerate(logs_qs, 1):
        row_num = idx + 4
        
        pihak = log.employee.name if log.employee else "-"
        notes = log.notes or "-"
        
        if not log.employee and log.notes:
            match = pattern.match(log.notes.strip())
            if match:
                name, nik, rest_of_notes = match.groups()
                pihak = name
                notes = rest_of_notes.strip()
                
        data = [
            idx,
            log.timestamp.strftime('%d %b %Y %H:%M'),
            log.get_action_display(),
            notes,
            pihak
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 2, 3]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    for col, width in zip(['A','B','C','D','E'], [5, 20, 15, 40, 20]):
        ws.column_dimensions[col].width = width
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Riwayat_{unit.unit_code}_{now_str}.xlsx"'
    wb.save(response)
    return response

@login_required
def asset_units_export_excel(request, asset_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.http import HttpResponse
    from django.shortcuts import get_object_or_404
    import datetime
    
    asset = get_object_or_404(Asset, id=asset_id)
    tab = request.GET.get('tab', 'semua')
    
    units = asset.units.all()
    if tab == 'tersedia':
        units = units.filter(status='tersedia')
    elif tab == 'digunakan':
        units = units.filter(status='digunakan')
    elif tab == 'rusak':
        units = units.filter(status='rusak')
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Daftar Unit"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:F1')
    ws['A1'] = "LAPORAN DAFTAR UNIT ASET"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    ws.merge_cells('A2:F2')
    ws['A2'] = f"{asset.name}"
    ws['A2'].alignment = center_align
    
    headers = ['No', 'Kode Unit', 'Area', 'Status', 'Lokasi/Pengguna', '-']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    for idx, u in enumerate(units, 1):
        row_num = idx + 4
        
        status_map = {'tersedia': 'Tersedia', 'digunakan': 'Digunakan', 'rusak': 'Maintenance'}
        lokasi = "-"
        
        if u.status == 'digunakan':
            assignment = AssetAssignment.objects.filter(unit=u, status='aktif').first()
            if assignment and assignment.employee:
                lokasi = assignment.employee.name
        elif u.status == 'rusak' and u.location:
            lokasi = u.location
            
        data = [
            idx,
            u.unit_code,
            u.area or '-',
            status_map.get(u.status, u.status),
            lokasi,
            u.location or '-'
        ]
        
        for col_num, cell_value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=cell_value)
            cell.border = border
            if col_num in [1, 2, 4]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    for col, width in zip(['A','B','C','D','E','F'], [5, 15, 20, 15, 30, 30]):
        ws.column_dimensions[col].width = width
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    now_str = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    response['Content-Disposition'] = f'attachment; filename="Unit_{asset.asset_code}_{now_str}.xlsx"'
    wb.save(response)
    return response

