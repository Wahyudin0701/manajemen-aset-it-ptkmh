import re
from datetime import datetime
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.db import transaction
from django.db.models import Count, Q
from django.core.paginator import Paginator

from assets.models import ConfigurationItem, CIHistory, AssetUnit, Employee, UnitLog
from assets.permissions import operator_required


@login_required
def ci_list(request):
    """Menampilkan katalog Configuration Item (CI) dengan kartu ringkasan,
    filter kategori, status operasional, kondisi fisik, dan paginasi 25 baris.
    """
    search_query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('status', '').strip()
    category_filter = request.GET.get('category', '').strip()
    condition_filter = request.GET.get('condition', '').strip()

    # Base queryset
    ci_qs = ConfigurationItem.objects.select_related(
        'installed_on', 'installed_on__asset', 'installed_on__assigned_to', 'assigned_to'
    )

    # Ringkasan KPI operasional
    counts = {
        'total': ConfigurationItem.objects.count(),
        'in_use': ConfigurationItem.objects.filter(status='in_use').count(),
        'available': ConfigurationItem.objects.filter(status='available').count(),
        'under_repair': ConfigurationItem.objects.filter(status__in=['under_inspection', 'under_repair']).count(),
        'retired': ConfigurationItem.objects.filter(status='retired').count(),
    }

    # Penerapan Filter
    if status_filter:
        if status_filter == 'repair':
            ci_qs = ci_qs.filter(status__in=['under_inspection', 'under_repair'])
        else:
            ci_qs = ci_qs.filter(status=status_filter)

    if category_filter:
        ci_qs = ci_qs.filter(category=category_filter)

    if condition_filter:
        ci_qs = ci_qs.filter(condition=condition_filter)

    if search_query:
        ci_qs = ci_qs.filter(
            Q(ci_code__icontains=search_query) |
            Q(name__icontains=search_query) |
            Q(brand__icontains=search_query) |
            Q(model_series__icontains=search_query) |
            Q(serial_number__icontains=search_query) |
            Q(specifications__icontains=search_query) |
            Q(installed_on__unit_code__icontains=search_query) |
            Q(assigned_to__name__icontains=search_query)
        )

    # Paginasi 25 item per halaman
    paginator = Paginator(ci_qs, 25)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Master data untuk modal aksi
    available_units = AssetUnit.objects.select_related('asset', 'assigned_to').filter(
        status__in=['tersedia', 'digunakan']
    ).order_by('unit_code')

    available_employees = Employee.objects.filter(is_active=True).order_by('name')

    # CI yang berstatus available untuk pilihan modal ganti/upgrade
    spare_cis = ConfigurationItem.objects.filter(status='available').order_by('name')

    return render(request, 'ci_list.html', {
        'items': page_obj,
        'page_obj': page_obj,
        'counts': counts,
        'search_query': search_query,
        'status_filter': status_filter,
        'category_filter': category_filter,
        'condition_filter': condition_filter,
        'categories': ConfigurationItem.CATEGORY_CHOICES,
        'status_choices': ConfigurationItem.STATUS_CHOICES,
        'condition_choices': ConfigurationItem.CONDITION_CHOICES,
        'available_units': available_units,
        'available_employees': available_employees,
        'spare_cis': spare_cis,
    })


@login_required
def ci_detail(request, ci_id):
    """Menampilkan informasi komprehensif suatu CI: spesifikasi, relasi aktif,
    dan timeline riwayat siklus hidup (audit trail).
    """
    ci = get_object_or_404(
        ConfigurationItem.objects.select_related(
            'installed_on', 'installed_on__asset', 'installed_on__assigned_to', 'assigned_to'
        ),
        id=ci_id
    )

    history_logs = ci.history.select_related(
        'related_unit', 'related_unit__asset', 'related_employee', 'performed_by'
    ).all()

    return render(request, 'ci_detail.html', {
        'ci': ci,
        'history_logs': history_logs,
    })


@login_required
@operator_required
def ci_create(request):
    """Menambahkan Configuration Item baru ke dalam sistem."""
    if request.method == 'POST':
        ci_code = request.POST.get('ci_code', '').strip().upper()
        name = request.POST.get('name', '').strip()
        category = request.POST.get('category', 'storage')
        brand = request.POST.get('brand', '').strip()
        model_series = request.POST.get('model_series', '').strip()
        serial_number = request.POST.get('serial_number', '').strip()
        specifications = request.POST.get('specifications', '').strip()
        condition = request.POST.get('condition', 'new')
        status = request.POST.get('status', 'available')
        location = request.POST.get('location', '').strip() or "Gudang IT"
        purchase_date = request.POST.get('purchase_date') or None
        warranty_expiry = request.POST.get('warranty_expiry') or None
        price = request.POST.get('price') or 0
        notes = request.POST.get('notes', '').strip()

        # Auto-generate CI code jika tidak diisi
        if not ci_code:
            prefix_map = {
                'storage': 'CI-STR',
                'memory': 'CI-MEM',
                'processor': 'CI-CPU',
                'motherboard': 'CI-MB',
                'gpu': 'CI-GPU',
                'peripheral': 'CI-PER',
                'network': 'CI-NET',
                'power': 'CI-PWR',
                'other': 'CI-GEN',
            }
            prefix = prefix_map.get(category, 'CI-IT')
            last_item = ConfigurationItem.objects.filter(ci_code__startswith=f"{prefix}-").order_by('-id').first()
            counter = 1
            if last_item:
                match = re.search(r'-(\d+)$', last_item.ci_code)
                if match:
                    counter = int(match.group(1)) + 1
            ci_code = f"{prefix}-{counter:03d}"

        # Validasi keunikan CI code
        if ConfigurationItem.objects.filter(ci_code=ci_code).exists():
            msg = f'Kode CI "{ci_code}" sudah terdaftar di sistem!'
            if request.headers.get('Accept') == 'application/json' or request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'status': 'error', 'message': msg}, status=400)
            messages.error(request, msg)
            return redirect('ci_list')

        ci = ConfigurationItem.objects.create(
            ci_code=ci_code,
            name=name,
            category=category,
            brand=brand,
            model_series=model_series,
            serial_number=serial_number,
            specifications=specifications,
            condition=condition,
            status=status,
            location=location,
            purchase_date=purchase_date,
            warranty_expiry=warranty_expiry,
            price=price,
            notes=notes
        )

        CIHistory.objects.create(
            ci=ci,
            action='created',
            performed_by=request.user,
            notes=f"Didaftarkan ke sistem dengan status {ci.get_status_display()}, kondisi {ci.get_condition_display()} di {location}."
        )

        messages.success(request, f'Configuration Item "{ci.ci_code} - {ci.name}" berhasil ditambahkan.')
        if request.headers.get('Accept') == 'application/json' or request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'status': 'success', 'ci_id': ci.id})

    return redirect('ci_list')


@login_required
@operator_required
def ci_edit(request, ci_id):
    """Memperbarui informasi metadata dan spesifikasi Configuration Item."""
    ci = get_object_or_404(ConfigurationItem, id=ci_id)

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        brand = request.POST.get('brand', '').strip()
        model_series = request.POST.get('model_series', '').strip()
        serial_number = request.POST.get('serial_number', '').strip()
        specifications = request.POST.get('specifications', '').strip()
        condition = request.POST.get('condition', ci.condition)
        location = request.POST.get('location', ci.location)
        warranty_expiry = request.POST.get('warranty_expiry') or None
        notes = request.POST.get('notes', '').strip()

        old_condition = ci.condition
        ci.name = name
        ci.brand = brand
        ci.model_series = model_series
        ci.serial_number = serial_number
        ci.specifications = specifications
        ci.condition = condition
        ci.location = location
        ci.warranty_expiry = warranty_expiry
        ci.notes = notes
        ci.save()

        if old_condition != condition:
            CIHistory.objects.create(
                ci=ci,
                action='status_change',
                performed_by=request.user,
                notes=f"Kondisi fisik diperbarui dari {dict(ConfigurationItem.CONDITION_CHOICES).get(old_condition)} ke {ci.get_condition_display()}."
            )

        messages.success(request, f'Perubahan pada {ci.ci_code} berhasil disimpan.')
        if request.headers.get('Accept') == 'application/json' or request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'status': 'success'})

    return redirect('ci_detail', ci_id=ci.id)


@login_required
@operator_required
def ci_assign(request, ci_id):
    """Alur 'Gunakan Untuk' (Assignment): Memasang komponen ke perangkat utama (AssetUnit)
    atau mengalokasikannya langsung ke Karyawan secara transaksional.
    """
    ci = get_object_or_404(ConfigurationItem, id=ci_id)

    if request.method == 'POST':
        usage_type = request.POST.get('usage_type')  # 'asset' atau 'employee'
        target_id = request.POST.get('target_id')
        notes = request.POST.get('notes', '').strip()

        if ci.status == 'in_use':
            messages.error(request, f'Komponen {ci.ci_code} sedang digunakan. Lepas terlebih dahulu sebelum dialokasikan.')
            return redirect('ci_list')

        if ci.status == 'retired':
            messages.error(request, f'Komponen {ci.ci_code} sudah dipensiunkan.')
            return redirect('ci_list')

        with transaction.atomic():
            if usage_type == 'asset':
                unit = get_object_or_404(AssetUnit, id=target_id)
                ci.installed_on = unit
                ci.assigned_to = None
                ci.status = 'in_use'
                ci.location = None
                ci.save()

                pic_info = f" (Pengguna: {unit.assigned_to.name})" if unit.assigned_to else ""
                log_text = f"Dipasang pada perangkat {unit.unit_code} ({unit.asset.name}){pic_info}. {notes}".strip()

                CIHistory.objects.create(
                    ci=ci,
                    action='installed',
                    related_unit=unit,
                    related_employee=unit.assigned_to,
                    performed_by=request.user,
                    notes=log_text
                )

                # Rekam juga di UnitLog perangkat induk
                UnitLog.objects.create(
                    unit=unit,
                    action='other',
                    employee=unit.assigned_to,
                    notes=f"Pemasangan komponen CI: {ci.name} ({ci.ci_code}). {notes}".strip()
                )

                messages.success(request, f'{ci.name} berhasil dipasang pada perangkat {unit.unit_code}.')

            elif usage_type == 'employee':
                employee = get_object_or_404(Employee, id=target_id, is_active=True)
                ci.assigned_to = employee
                ci.installed_on = None
                ci.status = 'in_use'
                ci.location = None
                ci.save()

                log_text = f"Dialokasikan langsung kepada karyawan {employee.name} ({employee.nik}). {notes}".strip()

                CIHistory.objects.create(
                    ci=ci,
                    action='assigned',
                    related_employee=employee,
                    performed_by=request.user,
                    notes=log_text
                )

                messages.success(request, f'{ci.name} berhasil dialokasikan kepada {employee.name}.')

        if request.headers.get('Accept') == 'application/json' or request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'status': 'success'})

    return redirect('ci_list')


@login_required
@operator_required
def ci_remove(request, ci_id):
    """Alur 'Lepas dari Perangkat / Pengguna': Memutus relasi aktif dan mengembalikan
    komponen ke gudang, pemeriksaan, atau perbaikan secara transaksional.
    """
    ci = get_object_or_404(ConfigurationItem, id=ci_id)

    if request.method == 'POST':
        new_status = request.POST.get('new_status', 'available')  # available, under_inspection, under_repair
        new_condition = request.POST.get('new_condition', ci.condition)
        location = request.POST.get('location', '').strip() or "Gudang IT"
        notes = request.POST.get('notes', '').strip()

        old_unit = ci.installed_on
        old_emp = ci.assigned_to

        with transaction.atomic():
            ci.installed_on = None
            ci.assigned_to = None
            ci.status = new_status
            ci.condition = new_condition
            ci.location = location
            ci.save()

            from_desc = ""
            if old_unit:
                from_desc = f"perangkat {old_unit.unit_code}"
            elif old_emp:
                from_desc = f"karyawan {old_emp.name}"

            log_text = f"Dilepas dari {from_desc}. Status kini {ci.get_status_display()}, kondisi {ci.get_condition_display()} di {location}. {notes}".strip()

            CIHistory.objects.create(
                ci=ci,
                action='removed',
                related_unit=old_unit,
                related_employee=old_emp,
                performed_by=request.user,
                notes=log_text
            )

            # Jika dilepas dari unit perangkat fisik, rekam di UnitLog
            if old_unit:
                UnitLog.objects.create(
                    unit=old_unit,
                    action='other',
                    employee=old_unit.assigned_to,
                    notes=f"Pelepasan komponen CI: {ci.name} ({ci.ci_code}). {notes}".strip()
                )

        messages.success(request, f'{ci.name} berhasil dilepas dan status diperbarui ke {ci.get_status_display()}.')
        if request.headers.get('Accept') == 'application/json' or request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'status': 'success'})

    return redirect('ci_list')


@login_required
@operator_required
def ci_replace(request, ci_id):
    """Alur 'Ganti / Upgrade Komponen': Menggantikan komponen lama yang terpasang pada
    suatu perangkat dengan komponen baru/cadangan dalam satu transaksi atomik.
    """
    old_ci = get_object_or_404(ConfigurationItem, id=ci_id)

    if request.method == 'POST':
        replacement_ci_id = request.POST.get('replacement_ci_id')
        reason = request.POST.get('reason', 'Replacement Kerusakan')
        old_ci_status = request.POST.get('old_ci_status', 'under_repair')  # under_repair, available, retired
        old_ci_condition = request.POST.get('old_ci_condition', 'damaged')
        old_ci_location = request.POST.get('old_ci_location', 'Gudang Perbaikan')
        notes = request.POST.get('notes', '').strip()

        if not old_ci.installed_on:
            messages.error(request, f'Komponen {old_ci.ci_code} tidak terpasang pada perangkat apapun.')
            return redirect('ci_list')

        new_ci = get_object_or_404(ConfigurationItem, id=replacement_ci_id)

        if new_ci.id == old_ci.id:
            messages.error(request, 'Komponen pengganti tidak boleh sama dengan komponen yang diganti.')
            return redirect('ci_list')

        if new_ci.status != 'available':
            messages.error(request, f'Komponen pengganti {new_ci.ci_code} tidak berstatus Tersedia (sedang {new_ci.get_status_display()}).')
            return redirect('ci_list')

        target_unit = old_ci.installed_on

        with transaction.atomic():
            # 1. Lepas komponen lama
            old_ci.installed_on = None
            old_ci.assigned_to = None
            old_ci.status = old_ci_status
            old_ci.condition = old_ci_condition
            old_ci.location = old_ci_location
            old_ci.save()

            # 2. Pasang komponen baru
            new_ci.installed_on = target_unit
            new_ci.assigned_to = None
            new_ci.status = 'in_use'
            new_ci.location = None
            new_ci.save()

            # 3. Catat history komponen lama
            CIHistory.objects.create(
                ci=old_ci,
                action='replaced',
                related_unit=target_unit,
                performed_by=request.user,
                notes=f"Dilepas dari unit {target_unit.unit_code} dan digantikan oleh {new_ci.ci_code} ({new_ci.name}) [Alasan: {reason}]. {notes}".strip()
            )

            # 4. Catat history komponen baru
            CIHistory.objects.create(
                ci=new_ci,
                action='installed',
                related_unit=target_unit,
                performed_by=request.user,
                notes=f"Dipasang pada unit {target_unit.unit_code} menggantikan komponen lama {old_ci.ci_code} ({old_ci.name}) [Alasan: {reason}]. {notes}".strip()
            )

            # 5. Rekam jejak di UnitLog perangkat induk
            UnitLog.objects.create(
                unit=target_unit,
                action='other',
                employee=target_unit.assigned_to,
                notes=f"Upgrade/Penggantian Komponen: {old_ci.name} digantikan oleh {new_ci.name}. Alasan: {reason}. {notes}".strip()
            )

        messages.success(request, f'Penggantian komponen pada unit {target_unit.unit_code} berhasil: {new_ci.ci_code} menggantikan {old_ci.ci_code}.')
        if request.headers.get('Accept') == 'application/json' or request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'status': 'success'})

    return redirect('ci_list')


@login_required
@operator_required
def ci_report_damage(request, ci_id):
    """Melaporkan keluhan kerusakan pada CI dan mengarahkan ke status pemeriksaan/perbaikan."""
    ci = get_object_or_404(ConfigurationItem, id=ci_id)

    if request.method == 'POST':
        action_type = request.POST.get('action_type', 'repair_in')  # 'repair_in' atau 'repair_out'
        condition = request.POST.get('condition', 'damaged')
        notes = request.POST.get('notes', '').strip()

        if action_type == 'repair_in':
            ci.status = 'under_repair'
            ci.condition = condition
            ci.save()

            CIHistory.objects.create(
                ci=ci,
                action='repair_in',
                performed_by=request.user,
                notes=f"Dilaporkan rusak / masuk perbaikan. Kondisi: {ci.get_condition_display()}. Catatan: {notes}"
            )
            messages.success(request, f'Komponen {ci.ci_code} ditandai masuk perbaikan.')

        elif action_type == 'repair_out':
            ci.status = 'available'
            ci.condition = 'good'
            ci.location = request.POST.get('location', 'Gudang IT')
            ci.save()

            CIHistory.objects.create(
                ci=ci,
                action='repair_out',
                performed_by=request.user,
                notes=f"Perbaikan selesai, kondisi dinyatakan bagus dan siap digunakan di {ci.location}. Catatan: {notes}"
            )
            messages.success(request, f'Komponen {ci.ci_code} selesai diperbaiki dan siap digunakan.')

    return redirect(request.META.get('HTTP_REFERER', 'ci_list'))


@login_required
def ci_export_excel(request):
    """Menghasilkan laporan Excel profesional untuk katalog Configuration Item."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    search_query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('status', '').strip()
    category_filter = request.GET.get('category', '').strip()
    condition_filter = request.GET.get('condition', '').strip()

    ci_qs = ConfigurationItem.objects.select_related(
        'installed_on', 'installed_on__asset', 'installed_on__assigned_to', 'assigned_to'
    ).order_by('category', 'ci_code')

    if status_filter:
        if status_filter == 'repair':
            ci_qs = ci_qs.filter(status__in=['under_inspection', 'under_repair'])
        else:
            ci_qs = ci_qs.filter(status=status_filter)

    if category_filter:
        ci_qs = ci_qs.filter(category=category_filter)

    if condition_filter:
        ci_qs = ci_qs.filter(condition=condition_filter)

    if search_query:
        ci_qs = ci_qs.filter(
            Q(ci_code__icontains=search_query) |
            Q(name__icontains=search_query) |
            Q(brand__icontains=search_query) |
            Q(model_series__icontains=search_query) |
            Q(serial_number__icontains=search_query) |
            Q(specifications__icontains=search_query) |
            Q(installed_on__unit_code__icontains=search_query) |
            Q(assigned_to__name__icontains=search_query)
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Katalog CI"

    title_font = Font(size=14, bold=True, color="002060")
    subtitle_font = Font(size=10, italic=True, color="555555")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    ws.merge_cells('A1:J1')
    ws['A1'] = "LAPORAN CONFIGURATION ITEM (CI) - PT. KERINCI HYDROPOWER"
    ws['A1'].font = title_font
    ws['A1'].alignment = left_align

    ws.merge_cells('A2:J2')
    ws['A2'] = f"Tanggal Ekspor: {datetime.now().strftime('%d %B %Y %H:%M')} WIB | Dicetak oleh: {request.user.get_full_name() or request.user.username}"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = left_align

    headers = [
        "No", "Kode CI", "Nama Komponen", "Kategori", "Merek & Model",
        "Nomor Seri (S/N)", "Status", "Kondisi", "Perangkat Terpasang", "PIC / Pengguna"
    ]
    ws.append([])
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for idx, ci in enumerate(ci_qs, 1):
        installed = f"{ci.installed_on.unit_code} ({ci.installed_on.asset.name})" if ci.installed_on else "-"
        pic = ci.assigned_to.name if ci.assigned_to else (ci.installed_on.assigned_to.name if ci.installed_on and ci.installed_on.assigned_to else "-")
        row = [
            idx,
            ci.ci_code,
            ci.name,
            ci.get_category_display(),
            f"{ci.brand or ''} {ci.model_series or ''}".strip() or "-",
            ci.serial_number or "-",
            ci.get_status_display(),
            ci.get_condition_display(),
            installed,
            pic
        ]
        ws.append(row)
        curr_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 2, 7, 8]:
                c.alignment = center_align
            else:
                c.alignment = left_align

    col_widths = {'A': 6, 'B': 18, 'C': 26, 'D': 16, 'E': 20, 'F': 20, 'G': 16, 'H': 14, 'I': 24, 'J': 22}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"laporan_configuration_item_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def ci_export_pdf(request):
    """Menghasilkan laporan PDF resmi untuk Configuration Item yang seragam dengan modul lainnya."""
    import os
    import datetime
    from io import BytesIO
    from PIL import Image

    from django.conf import settings
    from django.http import HttpResponse
    from django.db.models import Q

    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm, cm
    from reportlab.lib.colors import HexColor, white
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.platypus.flowables import HRFlowable
    from reportlab.platypus import Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

    search_query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('status', '').strip()
    category_filter = request.GET.get('category', '').strip()
    condition_filter = request.GET.get('condition', '').strip()

    ci_qs = ConfigurationItem.objects.select_related(
        'installed_on', 'installed_on__asset', 'installed_on__assigned_to', 'assigned_to'
    ).order_by('category', 'ci_code')

    status_labels = {
        'available': 'Tersedia',
        'in_use': 'Sedang Digunakan',
        'repair': 'Maintenance',
        'under_repair': 'Maintenance',
        'under_inspection': 'Maintenance',
        'retired': 'Pensiun / Bricked',
    }
    status_label = status_labels.get(status_filter, 'Semua Status')

    if status_filter:
        if status_filter == 'repair':
            ci_qs = ci_qs.filter(status__in=['under_inspection', 'under_repair'])
        else:
            ci_qs = ci_qs.filter(status=status_filter)

    cat_dict = dict(ConfigurationItem.CATEGORY_CHOICES)
    cat_label = cat_dict.get(category_filter, 'Semua Kategori') if category_filter else 'Semua Kategori'
    if category_filter:
        ci_qs = ci_qs.filter(category=category_filter)

    cond_dict = dict(ConfigurationItem.CONDITION_CHOICES)
    cond_label = cond_dict.get(condition_filter, '') if condition_filter else ''
    if condition_filter:
        ci_qs = ci_qs.filter(condition=condition_filter)

    if search_query:
        ci_qs = ci_qs.filter(
            Q(ci_code__icontains=search_query) |
            Q(name__icontains=search_query) |
            Q(brand__icontains=search_query) |
            Q(model_series__icontains=search_query) |
            Q(serial_number__icontains=search_query) |
            Q(specifications__icontains=search_query) |
            Q(installed_on__unit_code__icontains=search_query) |
            Q(assigned_to__name__icontains=search_query)
        )

    navy = HexColor('#002060')
    alt_row = HexColor('#F8F9FC')
    border_color = HexColor('#D0D5DD')

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=2 * cm
    )
    page_width = landscape(A4)[0] - 3 * cm

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CustomTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=16, textColor=white, leading=20)
    subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Normal'], fontName='Helvetica', fontSize=10, textColor=HexColor('#B0C4DE'), leading=14)
    company_style = ParagraphStyle('CompanyName', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, textColor=white, alignment=TA_RIGHT, leading=14)
    company_sub_style = ParagraphStyle('CompanySub', parent=styles['Normal'], fontName='Helvetica', fontSize=8, textColor=HexColor('#B0C4DE'), alignment=TA_RIGHT, leading=11)

    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=11)
    cell_bold_style = ParagraphStyle('CellBoldStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11)
    cell_small = ParagraphStyle('CellSmall', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, textColor=HexColor('#666666'), leading=9)
    header_cell = ParagraphStyle('HeaderCell', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, textColor=white, alignment=TA_CENTER, leading=11)
    cell_center = ParagraphStyle('CellCenter', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, alignment=TA_CENTER)
    info_style = ParagraphStyle('InfoStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=9, textColor=HexColor('#555555'))
    info_bold = ParagraphStyle('InfoBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, textColor=HexColor('#333333'))

    elements = []
    now = datetime.datetime.now()
    bulan_indo = {
        1: 'Januari', 2: 'Februari', 3: 'Maret', 4: 'April',
        5: 'Mei', 6: 'Juni', 7: 'Juli', 8: 'Agustus',
        9: 'September', 10: 'Oktober', 11: 'November', 12: 'Desember'
    }
    tanggal_str = f"{now.day} {bulan_indo[now.month]} {now.year}"

    # Top Logos
    try:
        it_logo_path = os.path.join(settings.BASE_DIR, 'static', 'img', 'it_all_role_logo.png')
        kerinci_logo_path = os.path.join(settings.BASE_DIR, 'static', 'img', 'kerinci_logo.png')
        desired_height = 1.3 * cm
        it_img = Image.open(it_logo_path)
        it_rl_img = RLImage(it_logo_path, width=desired_height * (it_img.width / float(it_img.height)), height=desired_height)
        ker_img = Image.open(kerinci_logo_path)
        ker_rl_img = RLImage(kerinci_logo_path, width=desired_height * (ker_img.width / float(ker_img.height)), height=desired_height)
        logo_table = Table([[ker_rl_img, it_rl_img]], colWidths=[page_width/2, page_width/2])
        logo_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (0, 0), 'LEFT'),
            ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('TOPPADDING', (0, 0), (-1, -1), 0)
        ]))
        elements.append(logo_table)
    except Exception:
        pass

    category_label = f"Status: {status_label} - Kategori: {cat_label}"

    # Navy Header Banner
    header_data = [[
        [Paragraph('Laporan Configuration Item', title_style), Paragraph(f'{category_label} - {tanggal_str}', subtitle_style)],
        [Paragraph('SISTEM MANAJEMEN ASET IT', company_style), Paragraph('Dicetak otomatis dari sistem', company_sub_style)]
    ]]
    header_table = Table(header_data, colWidths=[page_width * 0.65, page_width * 0.35])
    header_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), navy),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (0, 0), 20),
        ('RIGHTPADDING', (-1, -1), (-1, -1), 20),
        ('TOPPADDING', (0, 0), (-1, -1), 14),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 14),
        ('ROUNDEDCORNERS', [6, 6, 0, 0]),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 8 * mm))

    # Metadata Grid
    total_data = ci_qs.count()
    info_data = [
        [Paragraph('Kategori', info_style), Paragraph(f': {category_label}', info_bold), Paragraph('Total Data', info_style), Paragraph(f': {total_data} komponen', info_bold)]
    ]
    row_2 = [Paragraph('Tanggal Cetak', info_style), Paragraph(f': {tanggal_str}, {now.strftime("%H:%M")} WIB', info_bold)]
    filters_desc = []
    if cond_label: filters_desc.append(f"Kondisi: {cond_label}")
    if search_query: filters_desc.append(f"Cari: '{search_query}'")
    if filters_desc:
        row_2.extend([Paragraph('Filter', info_style), Paragraph(f': {", ".join(filters_desc)}', info_bold)])
    else:
        row_2.extend(['', ''])
    info_data.append(row_2)

    info_table = Table(info_data, colWidths=[page_width * 0.12, page_width * 0.38, page_width * 0.12, page_width * 0.38])
    info_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 0)
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 6 * mm))
    elements.append(HRFlowable(width="100%", thickness=1, color=border_color, spaceAfter=4 * mm))

    # Table Header & Data
    table_header = [
        Paragraph('NO', header_cell),
        Paragraph('KODE ITEM', header_cell),
        Paragraph('NAMA & KOMPONEN', header_cell),
        Paragraph('KATEGORI', header_cell),
        Paragraph('NOMOR SERI (S/N)', header_cell),
        Paragraph('PERANGKAT / PENGGUNA', header_cell),
        Paragraph('KONDISI', header_cell),
        Paragraph('STATUS', header_cell)
    ]
    col_widths = [
        page_width * 0.05,
        page_width * 0.14,
        page_width * 0.23,
        page_width * 0.12,
        page_width * 0.13,
        page_width * 0.17,
        page_width * 0.08,
        page_width * 0.08
    ]
    table_data = [table_header]

    for idx, ci in enumerate(ci_qs, 1):
        brand_model = f"{ci.brand or ''} {ci.model_series or ''}".strip()
        comp_info = [Paragraph(ci.name, cell_bold_style)]
        if brand_model:
            comp_info.append(Paragraph(brand_model, cell_small))

        if ci.installed_on:
            pic_name = ci.installed_on.assigned_to.name if ci.installed_on.assigned_to else "Gudang"
            holder_info = [
                Paragraph(ci.installed_on.unit_code, cell_bold_style),
                Paragraph(f"{ci.installed_on.asset.name} (PIC: {pic_name})", cell_small)
            ]
        elif ci.assigned_to:
            holder_info = [
                Paragraph(ci.assigned_to.name, cell_bold_style),
                Paragraph("Pegang Langsung", cell_small)
            ]
        else:
            holder_info = Paragraph("Gudang / Workshop", cell_style)

        row = [
            Paragraph(str(idx), cell_center),
            Paragraph(ci.ci_code, cell_bold_style),
            comp_info,
            Paragraph(ci.get_category_display(), cell_center),
            Paragraph(ci.serial_number or "-", cell_center),
            holder_info,
            Paragraph(ci.get_condition_display(), cell_center),
            Paragraph(ci.get_status_display(), cell_center)
        ]
        table_data.append(row)

    if total_data == 0:
        table_data.append([
            Paragraph("Belum ada data Configuration Item yang cocok dengan kriteria laporan.", cell_center),
            "", "", "", "", "", "", ""
        ])

    data_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table_style = [
        ('BACKGROUND', (0, 0), (-1, 0), navy),
        ('TEXTCOLOR', (0, 0), (-1, 0), white),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('LINEBELOW', (0, 0), (-1, 0), 1, navy),
        ('LINEBELOW', (0, 1), (-1, -2), 0.5, HexColor('#E5E7EB')),
        ('LINEBELOW', (0, -1), (-1, -1), 1, border_color),
    ]
    if total_data == 0:
        table_style.append(('SPAN', (0, 1), (-1, 1)))
        table_style.append(('ALIGN', (0, 1), (-1, 1), 'CENTER'))
        table_style.append(('TOPPADDING', (0, 1), (-1, 1), 16))
        table_style.append(('BOTTOMPADDING', (0, 1), (-1, 1), 16))
    else:
        for i in range(1, len(table_data)):
            if i % 2 == 0:
                table_style.append(('BACKGROUND', (0, i), (-1, i), alt_row))

    data_table.setStyle(TableStyle(table_style))
    elements.append(data_table)

    elements.append(Spacer(1, 5 * mm))
    elements.append(Paragraph(
        f"Menampilkan {total_data} komponen — Dokumen ini digenerate secara otomatis oleh Sistem Manajemen Aset IT pada {tanggal_str}.",
        ParagraphStyle('F', parent=styles['Normal'], fontName='Helvetica', fontSize=8, textColor=HexColor('#999999'))
    ))

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    filename = f"Laporan_Configuration_Item_{now.strftime('%Y%m%d')}.pdf"
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response